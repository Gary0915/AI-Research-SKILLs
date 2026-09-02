"""Deterministic semantic typography governance for new planner material."""

from __future__ import annotations

from hashlib import sha256
import json
from pathlib import Path
import unicodedata
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt


class TypographyGovernorError(ValueError):
    """Raised when a caller bypasses a semantic typography role."""


_ROLE_SPECS = (
    ("deck_title", 32, "bold", "foreground", "left", 2),
    ("section_title", 28, "bold", "foreground", "left", 2),
    ("slide_title", 26, "bold", "foreground", "left", 2),
    ("body", 18, "normal", "foreground", "left", 6),
    ("body_emphasis", 18, "bold", "accent_primary", "left", 4),
    ("caption", 12, "normal", "muted", "left", 2),
    ("citation", 10, "normal", "muted", "left", 2),
    ("metric_primary", 22, "bold", "focus", "center", 2),
    ("metric_secondary", 14, "normal", "measurement_reference", "center", 2),
    ("table_header", 13, "bold", "foreground", "center", 2),
    ("table_body", 12, "normal", "foreground", "left", 4),
    ("formula", 16, "normal", "foreground", "center", 3),
    ("figure_label", 13, "bold", "foreground", "left", 2),
    ("callout", 15, "bold", "focus", "left", 3),
    ("footer", 9, "normal", "muted", "left", 1),
    ("page_number", 9, "normal", "muted", "right", 1),
)


def _canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _hash(value: Any) -> str:
    return sha256(_canonical(value).encode("utf-8")).hexdigest()


def build_presentation_typography_profile(root: Path) -> dict[str, Any]:
    """Build system-owned hierarchy while retaining truthful font uncertainty."""
    root = Path(root).resolve()
    # The style artifact establishes a valid active theme identity, but CP3
    # explicitly leaves exact professor-font family evidence unresolved.
    style = json.loads((root / "thesis-deck-system/artifacts/phase3/visual-style-profile.json").read_text(encoding="utf-8"))
    if style.get("style_profile_id") != "VSP003":
        raise TypographyGovernorError("approved style profile VSP003 is required")
    roles = []
    for role, size, weight, color_role, alignment, max_lines in _ROLE_SPECS:
        roles.append({
            "role": role,
            "latin_font_family": "Aptos",
            "cjk_font_family": "Aptos",
            "fallback_family": "Aptos",
            "font_size_pt": size,
            "minimum_font_size_pt": max(9, size - 3),
            "weight": weight,
            "italic": False,
            "color_role": color_role,
            "alignment": alignment,
            "vertical_alignment": "middle",
            "line_spacing": 1.0,
            "paragraph_space_before_pt": 0,
            "paragraph_space_after_pt": 0,
            "text_box_margins_pt": {"left": 3, "right": 3, "top": 2, "bottom": 2},
            "maximum_line_count": max_lines,
            "fit_policy": "static_fit_no_uncontrolled_shrink",
            "evidence_status": "synthetic_system_owned",
        })
    profile = {
        "schema_version": "1.0.0",
        "typography_profile_id": "PTP-001",
        "version": "1.0.0",
        "style_profile_id": style["style_profile_id"],
        "roles": roles,
        "font_family_fidelity": "insufficient_evidence",
        "theme_binding_valid": True,
        "native_render_verified": False,
        "source_evidence_ids": ["VSP003"],
        "status": "partial_structural_calibration",
    }
    profile["typography_profile_sha256"] = _hash({key: value for key, value in profile.items() if key != "typography_profile_sha256"})
    return profile


def resolve_typography(profile: dict[str, Any], role: str, *, override: dict[str, Any] | None = None) -> dict[str, Any]:
    """Resolve a role; arbitrary font/color/size overrides are fail-closed."""
    if profile.get("typography_profile_id") != "PTP-001":
        raise TypographyGovernorError("unknown typography profile")
    if override:
        permitted = {"role"}
        if set(override) - permitted:
            raise TypographyGovernorError("uncontrolled typography override")
    resolved = next((item for item in profile.get("roles", []) if item["role"] == role), None)
    if resolved is None:
        raise TypographyGovernorError("unknown semantic typography role")
    return dict(resolved)


def validate_editable_text(profile: dict[str, Any], role: str, text: str) -> str:
    """Validate the textual contract without pretending local font rendering."""
    resolve_typography(profile, role)
    if not isinstance(text, str) or not text:
        raise TypographyGovernorError("editable text must be a nonempty Unicode string")
    normalized = unicodedata.normalize("NFC", text)
    if any(unicodedata.category(character).startswith("C") and character not in "\n\t" for character in normalized):
        raise TypographyGovernorError("editable text contains a disallowed control character")
    return normalized


_COLOR_ROLE_RGB = {
    "background": "FFFFFF", "foreground": "1F1F1F", "muted": "666666",
    "border": "B7B7B7", "caption_background": "F2F2F2", "accent_primary": "333333",
    "focus": "B00020", "warning": "B00020", "measurement_reference": "4F6D8A",
}


def apply_typography_to_shape(shape: Any, profile: dict[str, Any], role: str) -> dict[str, Any]:
    """Apply one governed semantic role to an existing editable text shape."""
    typography = resolve_typography(profile, role)
    if not getattr(shape, "has_text_frame", False):
        raise TypographyGovernorError("typography target is not an editable text shape")
    frame = shape.text_frame
    margins = typography["text_box_margins_pt"]
    frame.margin_left = Pt(margins["left"])
    frame.margin_right = Pt(margins["right"])
    frame.margin_top = Pt(margins["top"])
    frame.margin_bottom = Pt(margins["bottom"])
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[typography["alignment"]]
    color = RGBColor.from_string(_COLOR_ROLE_RGB[typography["color_role"]])
    for paragraph in frame.paragraphs:
        paragraph.alignment = alignment
        for run in paragraph.runs:
            run.font.name = typography["fallback_family"]
            run.font.size = Pt(typography["font_size_pt"])
            run.font.bold = typography["weight"] == "bold"
            run.font.italic = typography["italic"]
            run.font.color.rgb = color
    return {"role": role, "font_size_pt": typography["font_size_pt"], "color_role": typography["color_role"], "evidence_status": typography["evidence_status"]}


def write_presentation_typography_profile(root: Path, destination: Path | None = None) -> Path:
    root = Path(root).resolve()
    destination = Path(destination or root / "thesis-deck-system/artifacts/phase3")
    destination.mkdir(parents=True, exist_ok=True)
    path = destination / "presentation-typography-profile.json"
    path.write_text(json.dumps(build_presentation_typography_profile(root), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path
